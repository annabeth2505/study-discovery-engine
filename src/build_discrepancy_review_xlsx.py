"""
build_discrepancy_review_xlsx.py
Turn the discrepancy findings into a review-friendly Excel workbook for Sam.

Tabs:
  Summary            -- counts by type + how to read each
  Host mismatch      -- deposit host conflicts with the paper (highest-value)
  Gut sample issue   -- samples may not be gut (highest severity)
  Multi-host         -- deposit understates real host diversity
  Host missing       -- deposit host field blank/generic (metadata gap, low urgency)

Each row: accession, paper title, ENA host, the specific conflict, and the note.
"""

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

FLAGGED = "../results/discrepancies_flagged.tsv"
OUT = "../results/discrepancy_review_for_sam.xlsx"

HEAD_FILL = PatternFill("solid", fgColor="2F5496")
HEAD_FONT = Font(name="Arial", bold=True, color="FFFFFF")
BODY_FONT = Font(name="Arial")
# severity tints
SEV = {
    "Host mismatch":    "FCE4D6",   # orange-ish, review
    "Gut sample issue": "F8CBAD",   # deeper orange, serious
    "Multi-host":       "FFF2CC",   # yellow, valuable
    "Host missing":     "E2EFDA",   # green, low urgency
}


def sheet(wb, name, df, cols, tint):
    ws = wb.create_sheet(name[:31])
    for j, (key, label, width) in enumerate(cols, 1):
        c = ws.cell(1, j, label)
        c.font = HEAD_FONT; c.fill = HEAD_FILL
        c.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(j)].width = width
    fill = PatternFill("solid", fgColor=tint)
    for i, (_, row) in enumerate(df.iterrows(), 2):
        for j, (key, label, width) in enumerate(cols, 1):
            val = row.get(key, "")
            c = ws.cell(i, j, "" if pd.isna(val) else str(val))
            c.font = BODY_FONT
            c.alignment = Alignment(vertical="top", wrap_text=(width > 30))
            if j == 1:
                c.fill = fill
    ws.freeze_panes = "A2"
    return ws


def build():
    f = pd.read_csv(FLAGGED, sep="\t")

    def has(flag): return f[f["discrepancy_flags"].astype(str).str.contains(flag)]

    wb = Workbook()
    wb.remove(wb.active)

    # ---- Summary tab ----
    ws = wb.create_sheet("Summary")
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 12
    ws.column_dimensions["C"].width = 70
    rows = [
        ("Discrepancy type", "Count", "How to read it"),
        ("Host mismatch", len(has("host_mismatch")),
         "Deposit host conflicts with the paper. Mostly real (wrong host, microbe in host field). ~2 may be general-term false positives (e.g. 'ruminant')."),
        ("Gut sample issue", len(has("gut_sample_issue")),
         "Sample-level metadata suggests NON-gut samples. Highest severity — could mean the study doesn't belong in a gut catalog."),
        ("Multi-host", len(has("multihost_issue")),
         "Deposit understates real host diversity (host field names one species, samples span several). ~5 are mouse-model-of-human-disease false positives."),
        ("Host missing", len(has("host_missing")),
         "Deposit host field blank/generic while paper names a host. Metadata completeness gap — low urgency, not a catalog error."),
        ("", "", ""),
        ("TOTAL flagged", len(f), "Out of 560 paper-linked studies cross-checked (paper vs deposit)."),
    ]
    for i, r in enumerate(rows, 1):
        for j, v in enumerate(r, 1):
            c = ws.cell(i, j, v)
            if i == 1:
                c.font = HEAD_FONT; c.fill = HEAD_FILL
            else:
                c.font = Font(name="Arial", bold=(j == 1))
            c.alignment = Alignment(vertical="top", wrap_text=(j == 3))
    ws.freeze_panes = "A2"

    # ---- per-type tabs ----
    cols_mismatch = [("study_accession","Accession",16),("paper_title","Paper title",45),
                     ("host_species","ENA host",18),("discrepancy_host","Conflict (paper vs deposit)",45),
                     ("discrepancy_notes","Note",50)]
    cols_gut = [("study_accession","Accession",16),("paper_title","Paper title",45),
                ("host_species","ENA host",18),("discrepancy_gut","Gut-sample evidence",45),
                ("discrepancy_notes","Note",50)]
    cols_multi = [("study_accession","Accession",16),("paper_title","Paper title",45),
                  ("host_species","ENA host",18),("discrepancy_multihost","Multi-host evidence",45),
                  ("discrepancy_notes","Note",50)]
    cols_missing = [("study_accession","Accession",16),("paper_title","Paper title",45),
                    ("host_species","ENA host",18),("discrepancy_host_missing","Missing-host detail",35),
                    ("discrepancy_notes","Note",50)]

    sheet(wb, "Host mismatch", has("host_mismatch"), cols_mismatch, SEV["Host mismatch"])
    sheet(wb, "Gut sample issue", has("gut_sample_issue"), cols_gut, SEV["Gut sample issue"])
    sheet(wb, "Multi-host", has("multihost_issue"), cols_multi, SEV["Multi-host"])
    sheet(wb, "Host missing", has("host_missing"), cols_missing, SEV["Host missing"])

    wb.save(OUT)
    print(f"Flagged studies: {len(f)}")
    for t in ["host_mismatch","gut_sample_issue","multihost_issue","host_missing"]:
        print(f"  {t}: {len(has(t))}")
    print(f"Saved -> {OUT}")


if __name__ == "__main__":
    build()