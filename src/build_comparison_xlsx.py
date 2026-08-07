"""
build_comparison_xlsx.py
Final deliverable for Sam: a two-tab Excel comparison of the audited catalog
against his GMTOL approved set.

Tabs:
  "in common with gmtol"   -- studies in both his Sam_approve=Yes* set and mine
  "only Ananya's catalog"  -- studies only in mine (KEEP + REVIEW, with verdicts)

Both tabs use his GMTOL column order where mappable, plus the audit columns
(keep / is_human / llm_notes) so he sees each study's verdict and reasoning.
Excludes are dropped from the "only mine" tab (they're genuine rejections);
KEEP + REVIEW are included, with the verdict visible so he can filter.
"""

import re
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.utils.dataframe import dataframe_to_rows

CATALOG = "../results/yes_catalog.tsv"
AUDIT = "../results/catalog_audit.tsv"
GMTOL = "../data/GMTOL_MAG_STUDY_list_ananaya.tsv"
OUT = "../results/catalog_vs_gmtol_comparison.xlsx"

# columns from Sam's schema we can populate, in his order
SAM_COLS = ["study_accession", "paper_title", "first_author", "pub_year", "journal",
            "doi", "fulltext_link", "host_phylum", "host_class", "host_order",
            "host_family", "host_genus", "host_species", "n_host_species",
            "host_species_list", "country", "body_sites", "library_strategy_mix",
            "n_wgs_samples", "n_samples"]
AUDIT_COLS = ["keep", "is_human", "gut", "body_site", "host_verdict", "method_note", "llm_notes"]


def prj(cell):
    return re.findall(r'PRJ[A-Z]{2}\d+', str(cell))


def build():
    cat = pd.read_csv(CATALOG, sep="\t")
    audit = pd.read_csv(AUDIT, sep="\t")
    gmtol = pd.read_csv(GMTOL, sep="\t", dtype=str)

    # merge audit verdicts into the catalog (drop any overlapping cols first so
    # the audit's versions win cleanly and no _x/_y suffixes appear)
    overlap = [c for c in AUDIT_COLS if c in cat.columns and c != "study_accession"]
    if overlap:
        cat = cat.drop(columns=overlap)
    cat = cat.merge(audit[["study_accession"] + AUDIT_COLS], on="study_accession", how="left")

    # --- Sam's approved set (Yes*) and its accessions ---
    gmtol["_ap"] = gmtol["Sam_approve"].astype(str).str.strip().str.lower()
    approved = gmtol[gmtol["_ap"].str.startswith("yes")]
    approved_accs = set(a for cell in approved["accessions"].dropna() for a in prj(cell))
    # normalized-title recovery for alternate accessions
    def norm(t):
        return re.sub(r'\s+', ' ', re.sub(r'[^a-z0-9 ]', ' ', str(t).lower())).strip() if pd.notna(t) else ""
    approved_titles = {norm(t) for t in approved["title"].dropna()}

    cat["_nt"] = cat["paper_title"].apply(norm)
    cat["in_common"] = cat.apply(
        lambda r: bool((r["study_accession"] in approved_accs) or
                       (bool(r["_nt"]) and r["_nt"] in approved_titles)),
        axis=1)

    # ensure all Sam-schema columns exist (blank if we never had them)
    for c in SAM_COLS:
        if c not in cat.columns:
            cat[c] = ""

    keep_cols = SAM_COLS + AUDIT_COLS

    # GROUP 1: in common (any keep status, so he sees them all)
    common = cat[cat["in_common"]][keep_cols].copy()

    # GROUP 2: only mine, KEEP + REVIEW (drop EXCLUDE and drop the in-common ones)
    mine_only = cat[(~cat["in_common"]) & (cat["keep"].isin(["KEEP", "REVIEW"]))][keep_cols].copy()

    # sort each: KEEP before REVIEW, then by host_class
    order = {"KEEP": 0, "REVIEW": 1, "EXCLUDE": 2}
    for df in (common, mine_only):
        df["_o"] = df["keep"].map(order).fillna(3)
        df.sort_values(["_o", "host_class"], inplace=True)
        df.drop(columns="_o", inplace=True)

    print(f"In common with GMTOL: {len(common)}")
    print(f"Only Ananya's catalog (KEEP+REVIEW): {len(mine_only)}")
    print(f"  KEEP: {(mine_only['keep']=='KEEP').sum()}  REVIEW: {(mine_only['keep']=='REVIEW').sum()}")
    print(f"  flagged human in mine-only: {mine_only['is_human'].astype(str).str.lower().eq('true').sum()}")

    # ---- write xlsx ----
    wb = Workbook()
    wb.remove(wb.active)
    header_font = Font(name="Arial", bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="2F5496")
    review_fill = PatternFill("solid", fgColor="FFF2CC")   # light yellow for REVIEW rows
    human_fill = PatternFill("solid", fgColor="FCE4D6")    # light orange for human rows
    body_font = Font(name="Arial")

    def write_tab(name, df):
        ws = wb.create_sheet(name[:31])
        for j, col in enumerate(df.columns, 1):
            c = ws.cell(1, j, col)
            c.font = header_font
            c.fill = header_fill
            c.alignment = Alignment(vertical="center", wrap_text=True)
        for i, (_, row) in enumerate(df.iterrows(), 2):
            is_rev = str(row.get("keep")) == "REVIEW"
            is_hum = str(row.get("is_human")).lower() == "true"
            for j, col in enumerate(df.columns, 1):
                c = ws.cell(i, j, row[col])
                c.font = body_font
                c.alignment = Alignment(vertical="top", wrap_text=(col == "llm_notes"))
                if col == "keep" and is_rev:
                    c.fill = review_fill
                if col == "is_human" and is_hum:
                    c.fill = human_fill
        # column widths
        for j, col in enumerate(df.columns, 1):
            L = get_column_letter(j)
            if col == "llm_notes":
                ws.column_dimensions[L].width = 70
            elif col in ("paper_title", "host_species_list"):
                ws.column_dimensions[L].width = 40
            elif col in ("doi", "fulltext_link", "library_strategy_mix"):
                ws.column_dimensions[L].width = 22
            else:
                ws.column_dimensions[L].width = 15
        ws.freeze_panes = "A2"

    write_tab("in common with gmtol", common)
    write_tab("only Ananya's catalog", mine_only)

    wb.save(OUT)
    print(f"\nSaved -> {OUT}")


if __name__ == "__main__":
    build()