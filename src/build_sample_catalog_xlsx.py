"""
build_sample_catalog_xlsx.py
Turn the sample-level catalog into a readable Excel workbook.

Tabs:
  README        -- what this file is, column glossary, and the caveats that change
                   how the numbers should be read
  Samples       -- the full sample catalog, one row per biosample (frozen header +
                   autofilter, so it is browsable rather than a wall of text)
  Studies       -- one row per study: host pattern, completeness, diversity, seq mix
  Genomic block -- the 37 studies containing library_source=GENOMIC samples

Read-only presentation layer: it reformats what the TSVs already contain and changes
no values. Written with openpyxl in write_only mode -- 144k rows in normal mode is
slow and memory-hungry.
"""

import pandas as pd
from datetime import date
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.cell import WriteOnlyCell

R = "../results/"
SAMPLES = R + "samples.tsv"
PATTERNS = R + "study_sample_patterns.tsv"
GENOMIC = R + "genomic_block_profile.tsv"
OUT = R + "sample_catalog.xlsx"

HEAD_FILL = PatternFill("solid", fgColor="2F5496")
HEAD_FONT = Font(name="Arial", bold=True, color="FFFFFF")
BODY_FONT = Font(name="Arial")
TITLE_FONT = Font(name="Arial", bold=True, size=13)
NOTE_FONT = Font(name="Arial", italic=True, color="666666")

# column -> width, only where the default is unhelpful
WIDTHS = {
    "sample_accession": 18, "study_accession": 16, "host": 22,
    "host_scientific_name": 24, "host_tax_id": 12, "host_body_site": 18,
    "isolation_source": 24, "scientific_name": 26, "library_strategy": 17,
    "library_source": 20, "sample_title": 38, "country": 20,
    "host_resolved": 24, "host_resolved_from": 21, "n_runs": 9,
    "first_run_accession": 19, "all_run_accessions": 30,
    "distinct_host_list": 52, "seq_type_mix": 30, "sample_host_pattern": 21,
    "host_pattern_confirmed": 23, "host_pattern_llm_note": 60,
    "genomic_scientific_names": 46, "example_sample_titles": 38,
    "genomic_strategy_mix": 26, "study_seq_type_mix": 28,
}

GLOSSARY = [
    ("sample_accession", "Biosample ID (SAMN/SAMEA/SAMD). One row per sample -- this is the row key."),
    ("study_accession", "Parent study. Join key back to yes_catalog.tsv."),
    ("host", "Free-text host field as the submitter typed it (RAT, C57BL/6J, mice)."),
    ("host_scientific_name", "Structured host binomial, when the submitter filled it in."),
    ("host_tax_id", "NCBI taxonomy ID for the host. Collapses spelling variants to one species."),
    ("host_body_site", "Structured body site. Only 14% filled -- body site usually hides in sample_title."),
    ("isolation_source", "Free-text source (feces, cecal content)."),
    ("scientific_name", "The SAMPLE's organism -- usually a metagenome label, NOT the host."),
    ("library_strategy", "WGS / AMPLICON / RNA-Seq / WGA / Bisulfite-Seq / Hi-C."),
    ("library_source", "METAGENOMIC / GENOMIC / METATRANSCRIPTOMIC. Authoritative over strategy."),
    ("sample_title", "Submitter's free-text label. Often carries body site and experimental arm."),
    ("country", "Geographic origin."),
    ("host_resolved", "DERIVED: first non-generic of host_scientific_name -> host -> scientific_name."),
    ("host_resolved_from", "DERIVED: which field host_resolved came from, or UNRESOLVED. Audit trail."),
    ("n_runs", "How many sequencing runs collapsed into this sample row."),
    ("first_run_accession", "One representative run (the sorted-first), for reference."),
    ("all_run_accessions", "EVERY run accession for this sample, ';'-joined. Split to recover run grain."),
]

CAVEATS = [
    "Row grain is one BIOSAMPLE, not one run. ENA returns runs; 177,564 runs were "
    "collapsed into 144,665 samples. n_runs and all_run_accessions preserve the runs.",

    "'missing' is a LITERAL STRING that ENA returns, not a blank. Any fill-rate count "
    "must treat 'missing' as empty, or it overstates coverage.",

    "host_resolved is RESOLVED, not NORMALIZED. It can hold 'RAT', 'mice', 'C57BL/6J' "
    "(a strain) alongside clean binomials. Use host_resolved_from to see which field "
    "each value came from, and host_tax_id when you need a species-level count.",

    "host_resolved's scientific_name fallback sometimes picks up things that are NOT "
    "hosts -- cultured bacteria, 'negative control', sample types, per-animal IDs "
    "(see PRJEB26512, PRJNA1102860, PRJNA1184454). Treat scientific_name-derived "
    "hosts as provisional.",

    "Study counts: this file holds 1,047 studies, not the catalog's 1,044. Two entries "
    "(PRJEB43192, PRJNA1033630) are UMBRELLA projects whose data ENA registers under "
    "12 and 3 CHILD accessions; 5 of those children are not in yes_catalog.tsv on "
    "their own. See samples_fetch_status.tsv.",

    "6 rows in PRJNA63661 have a semicolon-joined LIST in sample_accession. That is how "
    "ENA returns those pooled/co-assembly runs -- it is upstream data, left unaltered "
    "rather than silently rewritten. They break the one-row-per-sample grain.",

    "GENOMIC samples (5,368 across 37 studies) are mostly NOT cultured isolates: 40% "
    "sequence the host animal's own DNA and 55% are metagenomes labelled GENOMIC. Only "
    "71 samples are genuine microbial isolates. See the 'Genomic block' tab.",

    "Nothing here is backfilled from the study-level catalog. Where a host is blank, "
    "the submitter left it blank -- the study-level host_species column often hides "
    "that by taking the mode across samples.",
]


def cell(ws, value, font=None, fill=None, wrap=False):
    """A styled cell for write_only sheets -- rows are streamed, so style on the way in."""
    c = WriteOnlyCell(ws, value=value)
    c.font = font or BODY_FONT
    if fill:
        c.fill = fill
    if wrap:
        c.alignment = Alignment(wrap_text=True, vertical="top")
    return c


def widths(ws, cols):
    for i, name in enumerate(cols, start=1):
        ws.column_dimensions[get_column_letter(i)].width = WIDTHS.get(name, 16)


def add_table(wb, title, df, freeze="A2"):
    """Write a DataFrame as a filterable sheet (write_only: header first, then rows)."""
    ws = wb.create_sheet(title)
    cols = list(df.columns)
    widths(ws, cols)                       # must precede the rows in write_only mode
    ws.freeze_panes = freeze
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{len(df) + 1}"

    ws.append([cell(ws, c, HEAD_FONT, HEAD_FILL) for c in cols])
    for row in df.itertuples(index=False, name=None):
        ws.append(["" if pd.isna(v) else v for v in row])
    return ws


def build_readme(wb, n_samples, n_studies, n_runs, pattern_counts):
    ws = wb.create_sheet("README")
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 108

    def line(a="", b="", font=BODY_FONT, bfont=None):
        ws.append([cell(ws, a, font),
                   cell(ws, b, bfont or BODY_FONT, wrap=True)])

    line("SAMPLE-LEVEL CATALOG", "", TITLE_FONT)
    line("", f"Generated {date.today().isoformat()} from ENA by src/fetch_samples.py")
    line()
    line("Samples (rows)", f"{n_samples:,}")
    line("Studies", f"{n_studies:,}")
    line("Sequencing runs behind them", f"{n_runs:,}")
    line()
    line("HOST PATTERN", "one row per study on the 'Studies' tab", TITLE_FONT)
    for k, v in pattern_counts.items():
        line(f"  {k}", f"{v:,} studies")
    line()
    line("TABS", "", TITLE_FONT)
    line("  Samples", "The full catalog, one row per biosample. Filter the header row.")
    line("  Studies", "One row per study: completeness, diversity, pattern label, seq mix.")
    line("  Genomic block", "The 37 studies containing library_source=GENOMIC samples.")
    line()
    line("COLUMNS", "", TITLE_FONT)
    for name, desc in GLOSSARY:
        line(f"  {name}", desc)
    line()
    line("READ THIS BEFORE COUNTING", "", TITLE_FONT)
    for c in CAVEATS:
        line("  -", c, BODY_FONT, NOTE_FONT)
    return ws


def run():
    print("reading TSVs...")
    samples = pd.read_csv(SAMPLES, sep="\t", dtype=str)
    patterns = pd.read_csv(PATTERNS, sep="\t")
    genomic = pd.read_csv(GENOMIC, sep="\t")

    n_runs = pd.to_numeric(samples["n_runs"], errors="coerce").fillna(0).sum()
    counts = patterns["host_pattern_confirmed"].value_counts().to_dict()

    # numeric columns should sort and filter as numbers, not text. Keep the original
    # where a value will not convert, rather than turning it into a silent NaN.
    for c in ["n_runs", "host_tax_id"]:
        conv = pd.to_numeric(samples[c], errors="coerce")
        samples[c] = conv.where(conv.notna() | samples[c].isna(), samples[c])

    wb = Workbook(write_only=True)
    build_readme(wb, len(samples), samples.study_accession.nunique(), int(n_runs), counts)
    print(f"writing Samples ({len(samples):,} rows)...")
    add_table(wb, "Samples", samples)
    print(f"writing Studies ({len(patterns):,} rows)...")
    add_table(wb, "Studies", patterns)
    print(f"writing Genomic block ({len(genomic):,} rows)...")
    add_table(wb, "Genomic block", genomic)

    print("saving...")
    wb.save(OUT)
    print(f"Saved -> {OUT}")


if __name__ == "__main__":
    run()
